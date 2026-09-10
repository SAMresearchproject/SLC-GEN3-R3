"""CR243 — Mass-Lift Channel Typing runner.

Reads frozen snapshot CR243_input_snapshot.xlsx (SHA-locked in precommit).
Applies the five precommitted typed forms with the locked tolerance discipline.
Runs all wrong controls. Computes unequal-pair residual quantization. Writes
all locked output files. Verdict per precommit gates.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import random
import re
from collections import Counter
from fractions import Fraction
from pathlib import Path
from typing import Any

import openpyxl

HERE = Path(__file__).parent
INPUT_XLSX = HERE / "CR243_input_snapshot.xlsx"
PRECOMMIT_MD = HERE / "CR243_PRECOMMIT.md"

# Locked SHA-256 of frozen inputs (verified at runner start)
EXPECTED_INPUT_SHA = "7f5d9cc4c20f4bb62c3e40a4825c793ccc020fed72abb44ec17fe30406d8e4e5"
EXPECTED_PRECOMMIT_SHA = "4ea9789806a8bd17e0c314cb6508ee9ab80c0212b49e68b0cdf94fd052d97dc2"

# Locked substrate atoms (read-only from CR238)
R = 12
D = 3
S = 8
ALPHA_H = 2
M_LEDGER = 126
L_LEDGER = 162


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    h.update(path.read_bytes())
    return h.hexdigest()


def classify(rc: str | None) -> str | None:
    if rc is None:
        return None
    if rc.startswith("color_triad"):
        return "color_triad"
    if rc.startswith("pair_write"):
        m = re.match(r"pair_write\[(\d+)\|anti(\d+)\]", rc)
        if not m:
            return "pair_write_unparsed"
        a, b = int(m.group(1)), int(m.group(2))
        if a == b:
            return "OCTET_pair" if a == 9 else "equal_bound_pair"
        return "unequal_pair"
    if (
        rc.startswith("plus_single_write")
        or rc.startswith("minus_single_write")
        or rc.startswith("neutral_single_write")
    ):
        return "single_write"
    if rc.startswith("hidden_source_support") or "support" in rc.lower():
        return "support"
    if "carrier" in rc.lower():
        return "carrier"
    return "unclassified"


def parse_pair(rc: str) -> tuple[int, int] | None:
    m = re.match(r"pair_write\[(\d+)\|anti(\d+)\]", rc or "")
    if not m:
        return None
    return int(m.group(1)), int(m.group(2))


def y_typed_form(cls: str, q_sign: str | None, q_abs: int | None,
                 R_: int = R, D_: int = D, S_: int = S) -> Fraction | None:
    """Locked typed form per channel class. Returns Fraction or None if N/A."""
    if cls == "color_triad":
        sign = -1 if q_sign == "negative" else (1 if q_sign == "positive" else 0)
        return sign * Fraction(q_abs + D_, R_)
    if cls == "equal_bound_pair":
        return Fraction(D_ + 2, R_ * (D_ + 1))
    if cls == "OCTET_pair":
        return Fraction(D_ * D_ + S_, D_ * S_ * S_)
    if cls == "single_write":
        return Fraction(0)
    if cls in ("support", "carrier"):
        return Fraction(1)
    return None


def ceil_tol(K_native: float) -> float:
    """Tolerance band = max(1e-5, 0.005 / K_native), per precommit."""
    if not K_native or K_native == 0:
        return 1e-5
    return max(1e-5, 0.005 / abs(K_native))


def load_rows() -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(INPUT_XLSX, data_only=True)
    ws = wb["Sheet1"]
    rows = []
    for r in range(2, ws.max_row + 1):
        rc = ws[f"D{r}"].value
        if rc is None:
            continue
        K = ws[f"K{r}"].value
        X = ws[f"X{r}"].value
        if K in (None, 0) or X is None:
            continue
        try:
            Y_obs = float(X) / float(K)
        except (TypeError, ValueError):
            continue
        rec = {
            "spreadsheet_row": r,
            "candidate_id": ws[f"B{r}"].value,
            "route_combination": rc,
            "K_native": float(K),
            "closure_depth": ws[f"L{r}"].value,
            "q_sign": ws[f"M{r}"].value,
            "q_abs": ws[f"N{r}"].value,
            "surface_sign": ws[f"U{r}"].value,
            "surface_depth": ws[f"V{r}"].value,
            "X_debit_stored": float(X),
            "Y_observed": Y_obs,
            "M_observed": ws[f"Y{r}"].value,
        }
        rec["class"] = classify(rc)
        pair = parse_pair(rc)
        if pair:
            rec["pair_a"], rec["pair_b"] = pair
        rows.append(rec)
    return rows


def evaluate_typed_forms(rows: list[dict[str, Any]],
                         R_: int = R, D_: int = D, S_: int = S
                         ) -> dict[str, Any]:
    """Per-class match counts. Returns summary plus per-row predictions."""
    per_row = []
    by_class = {}
    for rec in rows:
        cls = rec["class"]
        y_pred = y_typed_form(cls, rec["q_sign"], rec["q_abs"], R_, D_, S_)
        entry = {
            "spreadsheet_row": rec["spreadsheet_row"],
            "candidate_id": rec["candidate_id"],
            "route_combination": rec["route_combination"],
            "class": cls,
            "K_native": rec["K_native"],
            "Y_observed": rec["Y_observed"],
            "Y_predicted": float(y_pred) if y_pred is not None else None,
            "tolerance": ceil_tol(rec["K_native"]),
        }
        if y_pred is None:
            entry["match"] = None
            entry["abs_diff"] = None
        else:
            diff = abs(rec["Y_observed"] - float(y_pred))
            entry["abs_diff"] = diff
            entry["match"] = bool(diff <= entry["tolerance"])
        per_row.append(entry)
        by_class.setdefault(cls, {"total": 0, "match": 0, "evaluated": 0})
        by_class[cls]["total"] += 1
        if y_pred is not None:
            by_class[cls]["evaluated"] += 1
            if entry["match"]:
                by_class[cls]["match"] += 1
    return {"per_row": per_row, "by_class": by_class}


def wrong_control(rows: list[dict[str, Any]], label: str,
                  R_: int = R, D_: int = D, S_: int = S) -> dict[str, Any]:
    """Re-evaluate typed forms with one constant perturbed, return clean-class
    match counts.  Pass condition = degrades at least one clean class.
    """
    canonical = evaluate_typed_forms(rows)["by_class"]
    perturbed = evaluate_typed_forms(rows, R_, D_, S_)["by_class"]
    clean_classes = ("color_triad", "equal_bound_pair", "OCTET_pair",
                     "single_write", "support", "carrier")
    degrades = False
    detail = {}
    for cls in clean_classes:
        c = canonical.get(cls, {"match": 0, "evaluated": 0})
        p = perturbed.get(cls, {"match": 0, "evaluated": 0})
        detail[cls] = {
            "canonical_match": c["match"],
            "perturbed_match": p["match"],
            "evaluated": p["evaluated"],
            "degraded": p["match"] < c["match"],
        }
        if p["match"] < c["match"]:
            degrades = True
    return {
        "label": label,
        "perturbation": f"R={R_}, D={D_}, S={S_}",
        "degrades": degrades,
        "per_class": detail,
    }


def octet_specific_wc(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """WC-O1..O4: octet form under perturbed (D,S); WC-O5: |q|+S substitution."""
    octet_rows = [r for r in rows if r["class"] == "OCTET_pair"]
    results = []
    for label, D_, S_ in [("WC-O1", 2, 8), ("WC-O2", 4, 8),
                          ("WC-O3", 3, 7), ("WC-O4", 3, 9)]:
        y_pert = Fraction(D_ * D_ + S_, D_ * S_ * S_)
        per_row = []
        for rec in octet_rows:
            tol = ceil_tol(rec["K_native"])
            diff = abs(rec["Y_observed"] - float(y_pert))
            per_row.append({
                "spreadsheet_row": rec["spreadsheet_row"],
                "candidate_id": rec["candidate_id"],
                "Y_observed": rec["Y_observed"],
                "Y_perturbed": float(y_pert),
                "abs_diff": diff,
                "matches_perturbed": bool(diff <= tol),
            })
        results.append({
            "label": label,
            "perturbation": f"D={D_}, S={S_}, octet form (D^2+S)/(D*S^2)",
            "y_perturbed_value": float(y_pert),
            "y_perturbed_fraction": f"{y_pert.numerator}/{y_pert.denominator}",
            "breaks_canonical": not any(p["matches_perturbed"] for p in per_row),
            "per_row": per_row,
        })
    # WC-O5: |q|+S substitution
    per_row = []
    for rec in octet_rows:
        q_abs = rec["q_abs"] or 0
        y_sub = Fraction(q_abs + S, D * S * S)
        tol = ceil_tol(rec["K_native"])
        diff = abs(rec["Y_observed"] - float(y_sub))
        per_row.append({
            "spreadsheet_row": rec["spreadsheet_row"],
            "candidate_id": rec["candidate_id"],
            "Y_observed": rec["Y_observed"],
            "Y_substitution": float(y_sub),
            "abs_diff": diff,
            "matches_substitution": bool(diff <= tol),
            "note": "arithmetic agrees because |q|=9=D^2 on the canonical OCTET row",
        })
    results.append({
        "label": "WC-O5",
        "perturbation": "numerator substitution |q|+S in place of D^2+S",
        "typing_audit_flag": "weaker / row-dependent",
        "arithmetic_pass": all(p["matches_substitution"] for p in per_row),
        "per_row": per_row,
    })
    return results


def class_shuffle_wc(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Shuffle class labels then re-evaluate. Informational, not verdict-blocking."""
    rng = random.Random(20260623)
    shuffled = list(rows)
    classes = [r["class"] for r in shuffled]
    rng.shuffle(classes)
    relabeled = []
    for rec, new_cls in zip(shuffled, classes):
        rec2 = dict(rec)
        rec2["class"] = new_cls
        relabeled.append(rec2)
    canonical_by_cls = evaluate_typed_forms(rows)["by_class"]
    shuffled_by_cls = evaluate_typed_forms(relabeled)["by_class"]
    clean_classes = ("color_triad", "equal_bound_pair", "OCTET_pair",
                     "single_write", "support", "carrier")
    per_class = {}
    for cls in clean_classes:
        c = canonical_by_cls.get(cls, {"match": 0, "evaluated": 0})
        s = shuffled_by_cls.get(cls, {"match": 0, "evaluated": 0})
        canonical_rate = c["match"] / c["evaluated"] if c["evaluated"] else 0
        shuffled_rate = s["match"] / s["evaluated"] if s["evaluated"] else 0
        per_class[cls] = {
            "canonical_match_rate": canonical_rate,
            "shuffled_match_rate": shuffled_rate,
            "canonical_match": c["match"],
            "shuffled_match": s["match"],
            "shuffled_evaluated": s["evaluated"],
        }
    avg_canonical = sum(p["canonical_match_rate"] for p in per_class.values()) / len(per_class)
    avg_shuffled = sum(p["shuffled_match_rate"] for p in per_class.values()) / len(per_class)
    return {
        "label": "WC-CS",
        "perturbation": "class label shuffle, seed=20260623",
        "informational_only": True,
        "average_canonical_match_rate": avg_canonical,
        "average_shuffled_match_rate": avg_shuffled,
        "per_class": per_class,
    }


def unequal_pair_residuals(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Test quantization of Δ = Y_observed for unequal-pair rows against typed
    unit family.  Y_base_unequal = 0 (no precommitted base form).
    """
    units = {
        "1/R^2": Fraction(1, R * R),
        "1/(R*S)": Fraction(1, R * S),
        "1/(D*S^2)": Fraction(1, D * S * S),
        "1/R": Fraction(1, R),
        "1/M": Fraction(1, M_LEDGER),
        "1/L": Fraction(1, L_LEDGER),
    }
    tol = 1e-5
    unequal = [r for r in rows if r["class"] == "unequal_pair"]
    per_row = []
    counts_by_unit = Counter()
    matched_any = 0
    for rec in unequal:
        delta = rec["Y_observed"]
        row_tol = max(tol, ceil_tol(rec["K_native"]))
        best = None
        for name, u in units.items():
            n = round(delta / float(u))
            residual = delta - n * float(u)
            if abs(residual) <= row_tol:
                if best is None or abs(residual) < best["residual"]:
                    best = {
                        "unit": name,
                        "n": n,
                        "predicted": n * float(u),
                        "residual": abs(residual),
                    }
        if best:
            matched_any += 1
            counts_by_unit[best["unit"]] += 1
        per_row.append({
            "spreadsheet_row": rec["spreadsheet_row"],
            "route_combination": rec["route_combination"],
            "pair_a": rec.get("pair_a"),
            "pair_b": rec.get("pair_b"),
            "q_abs": rec["q_abs"],
            "surface_sign": rec["surface_sign"],
            "surface_depth": rec["surface_depth"],
            "K_native": rec["K_native"],
            "X_debit_stored": rec["X_debit_stored"],
            "delta": delta,
            "row_tolerance": row_tol,
            "best_unit": best["unit"] if best else None,
            "best_n": best["n"] if best else None,
            "best_predicted": best["predicted"] if best else None,
            "best_residual": best["residual"] if best else None,
            "quantized_match": best is not None,
        })
    return {
        "total_unequal_rows": len(unequal),
        "matched_any_unit": matched_any,
        "match_rate": matched_any / len(unequal) if unequal else 0,
        "matches_by_unit": dict(counts_by_unit),
        "per_row": per_row,
    }


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("")
        return
    fieldnames: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                fieldnames.append(k)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        for r in rows:
            w.writerow(r)


def verdict_from_gates(canonical_eval: dict[str, Any],
                       wcs: list[dict[str, Any]],
                       octet_wcs: list[dict[str, Any]],
                       residuals: dict[str, Any]) -> dict[str, Any]:
    by_class = canonical_eval["by_class"]

    def cls_all_match(cls: str) -> bool:
        c = by_class.get(cls, {"match": 0, "evaluated": 0})
        return c["evaluated"] > 0 and c["match"] == c["evaluated"]

    S1 = cls_all_match("color_triad")
    S2 = cls_all_match("equal_bound_pair")
    S3 = cls_all_match("OCTET_pair")
    S4 = cls_all_match("single_write")
    S5 = cls_all_match("support") and cls_all_match("carrier")

    wc_by_label = {w["label"]: w for w in wcs}
    S6 = all(wc_by_label[lbl]["degrades"] for lbl in ("WC-R1", "WC-R2", "WC-R3"))
    S7 = all(wc_by_label[lbl]["degrades"] for lbl in ("WC-D1", "WC-D2"))
    S8 = all(wc_by_label[lbl]["degrades"] for lbl in ("WC-S1", "WC-S2"))

    octet_break = {w["label"]: w for w in octet_wcs}
    S9 = all(octet_break[lbl]["breaks_canonical"]
             for lbl in ("WC-O1", "WC-O2", "WC-O3", "WC-O4"))

    strong_pass = all([S1, S2, S3, S4, S5, S6, S7, S8, S9])

    F1 = not all([S1, S2, S3, S4, S5])
    F2 = not all([S6, S7, S8])

    if F1 or F2:
        verdict = "FAIL"
        verdict_signature = "CR243_FAIL_TYPED_CHANNEL_HYPOTHESIS"
    elif strong_pass:
        verdict = "STRONG_PASS"
        verdict_signature = (
            "CR243_STRONG_PASS_MASSLIFT_IS_TYPED_CHANNEL_TABLE__"
            "FIVE_CLEAN_CLASSES_EXACT__SUBSTRATE_PERTURBATIONS_DEGRADE__"
            "OCTET_TYPING_LOCKED_AS_D2_PLUS_S_OVER_D_S2"
        )
    else:
        verdict = "BOUNDARY"
        verdict_signature = "CR243_BOUNDARY_CLEAN_CLASSES_EXACT_PARTIAL_WC_DEGRADATION"

    return {
        "verdict": verdict,
        "verdict_signature": verdict_signature,
        "strong_pass_conditions": {
            "S1_color_triad_all_match": S1,
            "S2_equal_bound_pair_all_match": S2,
            "S3_OCTET_match": S3,
            "S4_single_write_all_match": S4,
            "S5_support_and_carrier_all_match": S5,
            "S6_R_perturbations_degrade": S6,
            "S7_D_perturbations_degrade": S7,
            "S8_S_perturbations_degrade": S8,
            "S9_octet_specific_wcs_break_canonical": S9,
        },
        "fail_conditions": {
            "F1_any_clean_class_fails": F1,
            "F2_any_constant_WC_fails_to_degrade": F2,
        },
    }


def main() -> None:
    # Hash-lock check
    actual_input = sha256(INPUT_XLSX)
    actual_pre = sha256(PRECOMMIT_MD)
    if actual_input != EXPECTED_INPUT_SHA:
        raise SystemExit(f"INPUT SHA mismatch: got {actual_input}, expected {EXPECTED_INPUT_SHA}")
    if actual_pre != EXPECTED_PRECOMMIT_SHA:
        raise SystemExit(f"PRECOMMIT SHA mismatch: got {actual_pre}, expected {EXPECTED_PRECOMMIT_SHA}")

    rows = load_rows()

    # Canonical evaluation
    canonical_eval = evaluate_typed_forms(rows)

    # Substrate-constant wrong controls
    wcs = [
        wrong_control(rows, "WC-R1", R_=10),
        wrong_control(rows, "WC-R2", R_=11),
        wrong_control(rows, "WC-R3", R_=13),
        wrong_control(rows, "WC-D1", D_=2),
        wrong_control(rows, "WC-D2", D_=4),
        wrong_control(rows, "WC-S1", S_=7),
        wrong_control(rows, "WC-S2", S_=9),
    ]

    octet_wcs = octet_specific_wc(rows)
    shuffle_wc = class_shuffle_wc(rows)
    residuals = unequal_pair_residuals(rows)
    verdict = verdict_from_gates(canonical_eval, wcs, octet_wcs, residuals)

    # Inventory
    inventory = Counter(r["class"] for r in rows)

    # Class match summary
    class_match_summary = {}
    for cls, stats in canonical_eval["by_class"].items():
        class_match_summary[cls] = {
            "total": stats["total"],
            "evaluated": stats["evaluated"],
            "match": stats["match"],
            "match_rate": stats["match"] / stats["evaluated"] if stats["evaluated"] else None,
        }

    summary = {
        "verdict": verdict["verdict"],
        "verdict_signature": verdict["verdict_signature"],
        "input_sha256": EXPECTED_INPUT_SHA,
        "precommit_sha256": EXPECTED_PRECOMMIT_SHA,
        "substrate_atoms": {"R": R, "D": D, "S": S, "alpha_H": ALPHA_H,
                            "M": M_LEDGER, "L": L_LEDGER,
                            "kappa": "7117/768", "g": "1/64"},
        "row_inventory": dict(inventory),
        "class_match_summary": class_match_summary,
        "wrong_controls_constant_perturbation": [
            {k: v for k, v in w.items()} for w in wcs
        ],
        "wrong_controls_octet_specific": [
            {k: v for k, v in w.items() if k != "per_row"} for w in octet_wcs
        ],
        "class_shuffle_informational": {
            k: v for k, v in shuffle_wc.items() if k != "per_class" or True
        },
        "unequal_pair_residuals_summary": {
            "total_unequal_rows": residuals["total_unequal_rows"],
            "matched_any_unit": residuals["matched_any_unit"],
            "match_rate": residuals["match_rate"],
            "matches_by_unit": residuals["matches_by_unit"],
        },
        "strong_pass_conditions": verdict["strong_pass_conditions"],
        "fail_conditions": verdict["fail_conditions"],
    }

    (HERE / "CR243_summary.json").write_text(json.dumps(summary, indent=2, default=str))

    write_csv(HERE / "CR243_clean_class_predictions.csv", canonical_eval["per_row"])
    write_csv(HERE / "CR243_unequal_pair_residuals.csv", residuals["per_row"])

    wc_rows: list[dict[str, Any]] = []
    for w in wcs:
        for cls, d in w["per_class"].items():
            wc_rows.append({
                "label": w["label"],
                "perturbation": w["perturbation"],
                "class": cls,
                "canonical_match": d["canonical_match"],
                "perturbed_match": d["perturbed_match"],
                "evaluated": d["evaluated"],
                "degraded": d["degraded"],
            })
    for w in octet_wcs:
        if "per_row" not in w:
            continue
        for entry in w["per_row"]:
            wc_rows.append({
                "label": w["label"],
                "perturbation": w["perturbation"],
                "class": "OCTET_pair",
                "spreadsheet_row": entry["spreadsheet_row"],
                "Y_observed": entry["Y_observed"],
                "Y_perturbed_or_substitution": entry.get("Y_perturbed", entry.get("Y_substitution")),
                "abs_diff": entry["abs_diff"],
                "matches": entry.get("matches_perturbed", entry.get("matches_substitution")),
            })
    write_csv(HERE / "CR243_wrong_controls.csv", wc_rows)

    # Input manifest
    manifest = [
        {"file": "CR243_input_snapshot.xlsx", "sha256": EXPECTED_INPUT_SHA,
         "row_count": len(rows), "class_inventory": json.dumps(dict(inventory))},
        {"file": "CR243_PRECOMMIT.md", "sha256": EXPECTED_PRECOMMIT_SHA,
         "row_count": "", "class_inventory": ""},
    ]
    write_csv(HERE / "CR243_input_manifest.csv", manifest)

    print(f"VERDICT: {verdict['verdict']}")
    print(f"Signature: {verdict['verdict_signature']}")
    print("Class match summary:")
    for cls, stats in class_match_summary.items():
        print(f"  {cls:18s} {stats['match']:4d}/{stats['evaluated']:4d}  (total {stats['total']})")
    print(f"Unequal-pair residual quantization: {residuals['matched_any_unit']}/{residuals['total_unequal_rows']} = {residuals['match_rate']:.3f}")
    print(f"  Matches by unit: {dict(residuals['matches_by_unit'])}")
    print("WC results (constant perturbation):")
    for w in wcs:
        print(f"  {w['label']}  {w['perturbation']:30s}  degrades={w['degrades']}")
    print("OCTET-specific WCs:")
    for w in octet_wcs:
        flag = w.get("breaks_canonical", w.get("arithmetic_pass"))
        print(f"  {w['label']}  {w['perturbation'][:60]:60s}  breaks/note={flag}")
    print(f"Class-shuffle (informational): avg canonical {shuffle_wc['average_canonical_match_rate']:.3f}, avg shuffled {shuffle_wc['average_shuffled_match_rate']:.3f}")


if __name__ == "__main__":
    main()
