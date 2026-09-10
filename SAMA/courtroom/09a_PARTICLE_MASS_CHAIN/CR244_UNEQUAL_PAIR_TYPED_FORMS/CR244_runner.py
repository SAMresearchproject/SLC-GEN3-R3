"""CR244 — Unequal-Pair Typed Forms runner.

Locks the two unequal-pair lane forms (ORDINARY and OCTET-involved) from
CR244_PRECOMMIT.md and tests them against all 42 unequal-pair rows in the
SHA-locked source snapshot.  Runs all wrong controls per precommit.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import Counter
from fractions import Fraction
from pathlib import Path
from typing import Any

import openpyxl

HERE = Path(__file__).parent
INPUT_XLSX = HERE / "CR244_input_snapshot.xlsx"
PRECOMMIT_MD = HERE / "CR244_PRECOMMIT.md"

EXPECTED_INPUT_SHA = "7f5d9cc4c20f4bb62c3e40a4825c793ccc020fed72abb44ec17fe30406d8e4e5"
EXPECTED_PRECOMMIT_SHA = "1c1bd7c69ce8edf88cd8c1c164f822eefeee431adf0532ae8b141b610f9b6182"

# Locked substrate atoms
R = 12
D = 3
S = 8
M_LEDGER = 126


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    h.update(path.read_bytes())
    return h.hexdigest()


def y_ordinary(abs_diff: int, sign: int, R_: int = R, D_: int = D) -> Fraction:
    """Ordinary unequal-pair: sign(a-b) * (|a-b| + D) / R^4."""
    return Fraction(sign * (abs_diff + D_), R_ ** 4)


def y_octet(abs_diff: int, sign: int, R_: int = R, D_: int = D) -> Fraction:
    """OCTET-involved unequal-pair: sign(a-b) * (|a-b| + D^2/R) / R^4."""
    return Fraction(sign, 1) * Fraction(abs_diff * R_ + D_ * D_, R_ * R_ ** 4)


def load_unequal_rows() -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(INPUT_XLSX, data_only=True)
    ws = wb["Sheet1"]
    rows = []
    for r in range(2, ws.max_row + 1):
        rc = ws[f"D{r}"].value
        if not rc or not rc.startswith("pair_write"):
            continue
        m = re.match(r"pair_write\[(\d+)\|anti(\d+)\]", rc)
        if not m:
            continue
        a, b = int(m.group(1)), int(m.group(2))
        if a == b:
            continue
        K = ws[f"K{r}"].value
        X = ws[f"X{r}"].value
        if K in (None, 0) or X is None:
            continue
        sign = 1 if a > b else -1
        abs_diff = abs(a - b)
        octet_involved = (a == 9) or (b == 9)
        rows.append({
            "spreadsheet_row": r,
            "candidate_id": ws[f"B{r}"].value,
            "route_combination": rc,
            "a": a,
            "b": b,
            "abs_diff": abs_diff,
            "sign": sign,
            "K_native": float(K),
            "X_stored": float(X),
            "Y_observed": float(X) / float(K),
            "surface_sign": ws[f"U{r}"].value,
            "surface_depth": ws[f"V{r}"].value,
            "lane": "OCTET" if octet_involved else "ORDINARY",
        })
    return rows


def evaluate_lane(rows: list[dict[str, Any]],
                  R_: int = R, D_: int = D,
                  trigger_swap: bool = False,
                  denom_override: int | None = None,
                  ordinary_numer_offset: int | None = None,
                  ) -> dict[str, Any]:
    """Apply the two lane forms (with optional perturbations) to all rows.

    trigger_swap: if True, apply ORDINARY form to OCTET rows and vice versa.
    denom_override: if set, replace R^4 with this denominator.
    ordinary_numer_offset: if set, use (|a-b| + this) instead of (|a-b| + D)
        on the ORDINARY lane.
    """
    per_row = []
    counts = {"ORDINARY": {"total": 0, "x_round": 0, "y_tol": 0},
              "OCTET": {"total": 0, "x_round": 0, "y_tol": 0}}
    for rec in rows:
        lane = rec["lane"]
        effective_lane = lane
        if trigger_swap:
            effective_lane = "ORDINARY" if lane == "OCTET" else "OCTET"
        denom = denom_override if denom_override is not None else R_ ** 4
        if effective_lane == "ORDINARY":
            offset = ordinary_numer_offset if ordinary_numer_offset is not None else D_
            Y_pred = Fraction(rec["sign"] * (rec["abs_diff"] + offset), denom)
        else:
            Y_pred = Fraction(rec["sign"], 1) * Fraction(rec["abs_diff"] * R_ + D_ * D_, R_ * denom)
        Y_pred_float = float(Y_pred)
        X_pred_true = rec["K_native"] * Y_pred_float
        X_pred_rounded = round(X_pred_true, 2)
        match_x = (abs(X_pred_rounded - rec["X_stored"]) < 1e-9)
        tol = max(1e-5, 0.005 / abs(rec["K_native"]))
        match_y = abs(rec["Y_observed"] - Y_pred_float) <= tol
        counts[lane]["total"] += 1
        if match_x:
            counts[lane]["x_round"] += 1
        if match_y:
            counts[lane]["y_tol"] += 1
        per_row.append({
            "spreadsheet_row": rec["spreadsheet_row"],
            "route_combination": rec["route_combination"],
            "a": rec["a"],
            "b": rec["b"],
            "abs_diff": rec["abs_diff"],
            "lane": lane,
            "effective_lane": effective_lane,
            "K_native": rec["K_native"],
            "X_stored": rec["X_stored"],
            "Y_observed": rec["Y_observed"],
            "Y_predicted": Y_pred_float,
            "Y_predicted_fraction": f"{Y_pred.numerator}/{Y_pred.denominator}",
            "X_predicted_true": X_pred_true,
            "X_predicted_rounded": X_pred_rounded,
            "tolerance_Y": tol,
            "match_x_round": match_x,
            "match_y_tol": match_y,
        })
    return {"per_row": per_row, "counts": counts}


def wrong_control(rows: list[dict[str, Any]], label: str,
                  description: str, **kwargs) -> dict[str, Any]:
    canonical = evaluate_lane(rows)["counts"]
    perturbed = evaluate_lane(rows, **kwargs)["counts"]
    detail = {}
    degrades_any = False
    for lane in ("ORDINARY", "OCTET"):
        c = canonical[lane]["x_round"]
        p = perturbed[lane]["x_round"]
        detail[lane] = {
            "canonical_x_round": c,
            "perturbed_x_round": p,
            "total": canonical[lane]["total"],
            "degraded": p < c,
        }
        if p < c:
            degrades_any = True
    return {
        "label": label,
        "description": description,
        "kwargs": {k: (v if not isinstance(v, dict) else str(v)) for k, v in kwargs.items()},
        "degrades_any_lane": degrades_any,
        "per_lane": detail,
    }


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        path.write_text("")
        return
    fieldnames: list[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                fieldnames.append(k)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        for r in rows:
            w.writerow(r)


def main() -> None:
    actual_input = sha256(INPUT_XLSX)
    actual_pre = sha256(PRECOMMIT_MD)
    if actual_input != EXPECTED_INPUT_SHA:
        raise SystemExit(f"INPUT SHA mismatch: got {actual_input}, expected {EXPECTED_INPUT_SHA}")
    if actual_pre != EXPECTED_PRECOMMIT_SHA:
        raise SystemExit(f"PRECOMMIT SHA mismatch: got {actual_pre}, expected {EXPECTED_PRECOMMIT_SHA}")

    rows = load_unequal_rows()
    canonical_eval = evaluate_lane(rows)

    wcs = [
        wrong_control(rows, "WC-R1", "R = 10, D = 3", R_=10),
        wrong_control(rows, "WC-R2", "R = 11, D = 3", R_=11),
        wrong_control(rows, "WC-R3", "R = 13, D = 3", R_=13),
        wrong_control(rows, "WC-D1", "R = 12, D = 2", D_=2),
        wrong_control(rows, "WC-D2", "R = 12, D = 4", D_=4),
        wrong_control(rows, "WC-T1", "Trigger swap (ORDINARY <-> OCTET)", trigger_swap=True),
        wrong_control(rows, "WC-A1", "Alternate denominator R^3 = 1728", denom_override=R ** 3),
        wrong_control(rows, "WC-A2", "Alternate denominator R^2 * M = 18144", denom_override=R * R * M_LEDGER),
        wrong_control(rows, "WC-A3", "ORDINARY numerator (|a-b| + D^2) instead of (|a-b| + D)", ordinary_numer_offset=D * D),
    ]

    # Verdict gates
    ord_total = canonical_eval["counts"]["ORDINARY"]["total"]
    ord_match = canonical_eval["counts"]["ORDINARY"]["x_round"]
    oct_total = canonical_eval["counts"]["OCTET"]["total"]
    oct_match = canonical_eval["counts"]["OCTET"]["x_round"]

    S1 = (ord_match == ord_total)
    S2 = (oct_match == oct_total)
    wc_by_label = {w["label"]: w for w in wcs}
    S3 = all(wc_by_label[lbl]["per_lane"]["ORDINARY"]["degraded"] and
             wc_by_label[lbl]["per_lane"]["OCTET"]["degraded"]
             for lbl in ("WC-R1", "WC-R2", "WC-R3"))
    S4 = all(wc_by_label[lbl]["degrades_any_lane"] for lbl in ("WC-D1", "WC-D2"))
    S5 = (wc_by_label["WC-T1"]["per_lane"]["ORDINARY"]["degraded"] and
          wc_by_label["WC-T1"]["per_lane"]["OCTET"]["degraded"])
    S6 = all(wc_by_label[lbl]["degrades_any_lane"] for lbl in ("WC-A1", "WC-A2", "WC-A3"))

    strong_pass = all([S1, S2, S3, S4, S5, S6])
    F1 = not (S1 and S2)
    F2 = not (S3 and S4)
    F3 = not S5

    if F1 or F2 or F3:
        verdict = "FAIL"
        signature = "CR244_FAIL_UNEQUAL_PAIR_TYPED_FORMS"
    elif strong_pass:
        verdict = "STRONG_PASS"
        signature = (
            "CR244_STRONG_PASS_UNEQUAL_PAIR_TYPED_FORMS_CLOSED__"
            "ORDINARY_30_OF_30_AND_OCTET_12_OF_12_EXACT_AT_2DEC_X_STORAGE__"
            "R4_DENOMINATOR_LOCKED__D2_OVER_R_OCTET_CORRECTION_LOCKED__"
            "ALL_WRONG_CONTROLS_DEGRADE_AS_PREDICTED"
        )
    else:
        verdict = "BOUNDARY"
        signature = "CR244_BOUNDARY_FORMS_MATCH_PARTIAL_WC_DEGRADATION"

    summary = {
        "verdict": verdict,
        "verdict_signature": signature,
        "input_sha256": EXPECTED_INPUT_SHA,
        "precommit_sha256": EXPECTED_PRECOMMIT_SHA,
        "substrate_atoms": {"R": R, "D": D, "S": S, "R^4": R ** 4},
        "lane_forms": {
            "ORDINARY": "sign(a-b) * (|a-b| + D) / R^4",
            "OCTET": "sign(a-b) * (|a-b| + D^2/R) / R^4",
            "OCTET_trigger": "a == 9 or b == 9",
        },
        "canonical_counts": canonical_eval["counts"],
        "wrong_controls": wcs,
        "strong_pass_conditions": {
            "S1_ordinary_all_match": S1,
            "S2_octet_all_match": S2,
            "S3_R_perturbations_degrade_both": S3,
            "S4_D_perturbations_degrade_any": S4,
            "S5_trigger_swap_degrades_both": S5,
            "S6_alternate_forms_degrade": S6,
        },
        "fail_conditions": {
            "F1_lane_form_fails": F1,
            "F2_substrate_constant_does_not_degrade": F2,
            "F3_trigger_swap_does_not_degrade": F3,
        },
    }
    (HERE / "CR244_summary.json").write_text(json.dumps(summary, indent=2, default=str))

    write_csv(HERE / "CR244_unequal_pair_predictions.csv", canonical_eval["per_row"])

    wc_rows = []
    for w in wcs:
        for lane, d in w["per_lane"].items():
            wc_rows.append({
                "label": w["label"],
                "description": w["description"],
                "lane": lane,
                "canonical_x_round_match": d["canonical_x_round"],
                "perturbed_x_round_match": d["perturbed_x_round"],
                "total": d["total"],
                "degraded": d["degraded"],
            })
    write_csv(HERE / "CR244_wrong_controls.csv", wc_rows)

    inventory = Counter(r["lane"] for r in rows)
    manifest = [
        {"file": "CR244_input_snapshot.xlsx", "sha256": EXPECTED_INPUT_SHA,
         "row_count": len(rows), "lane_inventory": json.dumps(dict(inventory))},
        {"file": "CR244_PRECOMMIT.md", "sha256": EXPECTED_PRECOMMIT_SHA,
         "row_count": "", "lane_inventory": ""},
    ]
    write_csv(HERE / "CR244_input_manifest.csv", manifest)

    print(f"VERDICT: {verdict}")
    print(f"Signature: {signature}")
    print(f"ORDINARY:  {ord_match}/{ord_total} X-round match")
    print(f"OCTET:     {oct_match}/{oct_total} X-round match")
    print()
    print("WC results (canonical -> perturbed, X-round match):")
    for w in wcs:
        ord_d = w["per_lane"]["ORDINARY"]
        oct_d = w["per_lane"]["OCTET"]
        print(f"  {w['label']:5s}  {w['description']:50s}  ORD {ord_d['canonical_x_round']:2d}->{ord_d['perturbed_x_round']:2d}  OCT {oct_d['canonical_x_round']:2d}->{oct_d['perturbed_x_round']:2d}")
    print()
    print(f"Strong-pass: S1={S1} S2={S2} S3={S3} S4={S4} S5={S5} S6={S6}")


if __name__ == "__main__":
    main()
